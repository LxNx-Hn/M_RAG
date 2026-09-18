#!/usr/bin/env python3
"""Build a rich HTML staging document for one-time HWP Web transfer.

The canonical manuscript remains the Markdown source.  This helper expands its
Excel table insertion markers and its local figure insertion markers so that a
browser selection can be pasted into HWP Web with paragraph and table structure.
It does not change scores, citations, or source artifacts.
"""

from __future__ import annotations

import argparse
import html
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "FINALDOCS" / "MANUSCRIPT" / "GRADUATION_REPORT_TRANSFER_KO_60Q.md"
WORKBOOK = ROOT / "FINALDOCS" / "TABLES" / "TABLES_60Q.xlsx"
STYLE_MARKER = re.compile(r"\s*\[스타일\s*=\s*[^\]]+\]")
TABLE_MARKER = re.compile(r"^\[표삽입: .*?\| Sheet=([^|\]]+)\|.*\]$")
FIGURE_MARKER = re.compile(r"^\[그림삽입: ([^|\]]+)\|.*\]$")


def esc(value: object) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def table_html(workbook, sheet_name: str) -> str:
    ws = workbook[sheet_name.strip()]
    rows = list(ws.iter_rows(values_only=True))
    nonempty = [row for row in rows if any(value not in (None, "") for value in row)]
    # The workbook preserves a title and source row for HWP manual transfer.
    # The manuscript already places the caption above the marker, so importing
    # these rows again would duplicate the caption and source within the table.
    while nonempty and str(nonempty[0][0] or "").startswith("표 "):
        nonempty.pop(0)
    while nonempty and str(nonempty[0][0] or "").startswith("출처:"):
        nonempty.pop(0)
    if not nonempty:
        raise ValueError(f"Workbook sheet has no data: {sheet_name}")
    header, *body = nonempty
    header_html = "".join(f"<th>{esc(value)}</th>" for value in header)
    body_html = "".join(
        "<tr>" + "".join(f"<td>{esc(value)}</td>" for value in row) + "</tr>"
        for row in body
    )
    return f'<table class="thesis-table"><thead><tr>{header_html}</tr></thead><tbody>{body_html}</tbody></table>'


def tsv_table_html(lines: list[str]) -> str:
    rows = [line.split("\t") for line in lines if line.strip()]
    if not rows:
        raise ValueError("Expected copied-table rows")
    header, *body = rows
    header_html = "".join(f"<th>{esc(value)}</th>" for value in header)
    body_html = "".join(
        "<tr>" + "".join(f"<td>{esc(value)}</td>" for value in row) + "</tr>"
        for row in body
    )
    return f'<table class="thesis-table"><thead><tr>{header_html}</tr></thead><tbody>{body_html}</tbody></table>'


def cover_html() -> str:
    return """<div class="cover">
<h1>졸업자격실험보고서</h1>
<h2>한국어 질의 기반 영어 학술·기술 문서 RAG에서의 HyDE·CAD·SCD 조합 실험</h2>
<p>지도교수 김진석</p><p>컴퓨터공학과</p><p>동국대학교 WISE캠퍼스</p><p>문종건</p><p>2026</p>
</div>
<div class="approval">
<h1>졸업자격실험보고서</h1>
<p>논문제목</p>
<h2>Combination Experiments of HyDE, CAD, and SCD in RAG over English Academic Documents with Korean Queries</h2>
<p>문종건</p><p>지도교수 김진석</p>
<p>본 보고서를 졸업자격 실험보고서로 제출함.</p><p>2026년 10월 00일</p>
<p>문종건의 졸업자격 실험보고 통과를 인준함.</p><p>2026년 00월 00일</p>
<p>주 심 (인)</p><p>부 심 (인)</p><p>동국대학교 컴퓨터공학과</p>
</div>"""


def build_html() -> str:
    source = MANUSCRIPT.read_text(encoding="utf-8")
    start = source.index("# 국문초록")
    source = source[start:]
    workbook = load_workbook(WORKBOOK, data_only=False)
    body: list[str] = [cover_html()]
    in_code = False
    # 0=no copied table block, 1=waiting for its opening fence, 2=inside it.
    table_copy_mode = 0
    captured_table_lines: list[str] = []
    capture_table_copy = False
    inserted_table_marker = False
    for raw in source.splitlines():
        line = raw.strip()
        if line.startswith("```"):
            if table_copy_mode == 1:
                table_copy_mode = 2
                continue
            if table_copy_mode == 2:
                if capture_table_copy:
                    body.append(tsv_table_html(captured_table_lines))
                table_copy_mode = 0
                captured_table_lines = []
                capture_table_copy = False
                continue
            in_code = not in_code
            continue
        if line.startswith("[한글 표 복붙용"):
            table_copy_mode = 1
            capture_table_copy = not inserted_table_marker
            inserted_table_marker = False
            continue
        if table_copy_mode:
            if table_copy_mode == 2 and capture_table_copy:
                captured_table_lines.append(raw)
            continue
        if not line:
            continue
        table = TABLE_MARKER.match(line)
        if table:
            body.append(table_html(workbook, table.group(1)))
            inserted_table_marker = True
            continue
        figure = FIGURE_MARKER.match(line)
        if figure:
            source_path = Path(figure.group(1).replace("\\", "/"))
            try:
                relative = source_path.relative_to(Path("FINALDOCS/MANUSCRIPT"))
            except ValueError:
                relative = Path("..") / source_path.relative_to("FINALDOCS")
            body.append(
                f'<figure><img src="{esc(relative.as_posix())}" alt="논문 그림" />'
                "<figcaption>그림 파일 삽입 위치</figcaption></figure>"
            )
            continue
        cleaned = STYLE_MARKER.sub("", line).strip()
        if cleaned.startswith("# "):
            body.append(f'<h1 class="chapter">{esc(cleaned[2:])}</h1>')
        elif cleaned.startswith("## "):
            body.append(f'<h2>{esc(cleaned[3:])}</h2>')
        elif cleaned.startswith("### "):
            body.append(f'<h3>{esc(cleaned[4:])}</h3>')
        elif cleaned.startswith("[한글 수식 입력기"):
            body.append(f'<p class="equation-note">{esc(cleaned)}</p>')
        elif in_code:
            body.append(f'<p class="equation">{esc(cleaned)}</p>')
        else:
            body.append(f'<p>{esc(cleaned)}</p>')
        if cleaned:
            inserted_table_marker = False
    return """<!doctype html><html lang="ko"><head><meta charset="utf-8"><style>
body { font-family: "함초롬바탕", "Malgun Gothic", serif; font-size: 10pt; line-height: 1.6; margin: 0; }
h1 { font-size: 16pt; margin-top: 28pt; page-break-before: always; } h1:first-child { page-break-before: auto; }
h2 { font-size: 13pt; margin-top: 18pt; } h3 { font-size: 11pt; margin-top: 14pt; }
p { margin: 0 0 8pt; text-align: justify; } .cover, .approval { text-align: center; min-height: 920px; padding-top: 80px; } .cover p, .approval p { text-align: center; margin: 18pt 0; }
.thesis-table { border-collapse: collapse; width: 100%; font-size: 8.5pt; margin: 6pt 0 14pt; } .thesis-table th, .thesis-table td { border: 1px solid #555; padding: 4px; vertical-align: top; } .thesis-table th { background: #e9e9e9; }
figure { margin: 12pt auto; text-align: center; } figure img { max-width: 90%; height: auto; } figcaption { font-size: 9pt; margin-top: 5pt; }
.equation-note { font-size: 9pt; color: #555; } .equation { font-family: Cambria Math, serif; text-align: center; }
</style></head><body>""" + "\n".join(body) + "</body></html>"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    args.output.write_text(build_html(), encoding="utf-8")
    print(args.output.resolve())


if __name__ == "__main__":
    main()
