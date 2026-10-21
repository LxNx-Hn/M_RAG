"""Independent raw-score, pairing, question-group and submission-table checks."""

from __future__ import annotations
import json
import re
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "FINALDOCS/ANALYSIS"))
from analyze_questions import (  # noqa: E402
    BASE,
    CONDITIONS,
    GROUP_SIZES,
    METRICS,
    QUALITY,
    TYPES,
    comparisons,
    run_records,
)

FORBIDDEN = re.compile(
    r"ANOVA|분산분석|t-검정|t-test|p-value|p값|유의확률|통계적 유의성|검정력|Cohen|Holm|Wilcoxon|중앙값|\bmedian\b|\bDelta\b|델타|단정하지|보장하지|의미하지|확정하지",
    re.I,
)


def original_scores():
    result = {}
    for p in [
        ROOT
        / "experiments/results/evaluation/main-hyde-cad-scd-reference-scd-gpt4o-official/merged.ragas_scores.json",
        ROOT / "experiments/results/evaluation/ext60_gpt4o/merged.ragas_scores.json",
    ]:
        for r in json.loads(p.read_text(encoding="utf-8"))["per_sample"]:
            group = r["group"]
            code = f"H{int('hyde_on' in group)}C{int('cad' in group)}S{int('scd' in group)}"
            assert (r["query_id"], code) not in result
            result[r["query_id"], code] = r
    return result


def validate_observations(obs):
    assert (
        len(obs) == 480 and obs.query_id.nunique() == 60
    ), "60 x 8 observations required"
    assert not obs.duplicated(
        ["query_id", "condition"]
    ).any(), "Duplicate query-condition"
    assert (
        obs.groupby("query_id").condition.nunique() == 8
    ).all(), "A condition is missing"
    scores = original_scores()
    records = run_records()
    for r in obs.itertuples():
        original = scores[r.query_id, r.condition]
        for m in QUALITY:
            v = original[m]
            actual = getattr(r, m)
            assert (v is None and pd.isna(actual)) or (
                v is not None and np.isclose(v, actual, atol=1e-12, rtol=0)
            ), f"Original score changed: {r.query_id} {r.condition} {m}"
        assert (
            abs(
                r.duration_seconds
                - records[r.query_id, r.condition]["duration_seconds"]
            )
            < 1e-10
        )
        answer = records[r.query_id, r.condition]["generated_answer"]
        hangul = sum(
            0xAC00 <= ord(c) <= 0xD7A3
            or 0x1100 <= ord(c) <= 0x11FF
            or 0x3130 <= ord(c) <= 0x318F
            for c in answer
        )
        latin = sum(c.isascii() and c.isalpha() for c in answer)
        if hangul + latin:
            assert (
                abs(r.korean_ratio - hangul / (hangul + latin)) < 1e-12
            ), "Original Korean character ratio changed"
        else:
            # The preserved measurement has one digits-only output recorded as 0.
            assert (
                r.korean_ratio == 0 and answer.strip()
            ), "Digits-only original measurement differs"
    assert int(obs[QUALITY].isna().sum().sum()) == 5


def validate_summary(obs, summary, annotations):
    raw = obs.merge(
        annotations[
            [
                "query_id",
                "question_group",
                "question_type",
                "document",
                "answer_form",
                "information_units",
            ]
        ],
        on="query_id",
        validate="many_to_one",
    )
    for r in summary.itertuples():
        assert r.grouping in [
            "전체",
            "question_group",
            "question_type",
            "document",
            "answer_form",
            "information_units",
        ]
        g = raw if r.grouping == "전체" else raw[raw[r.grouping].astype(str) == r.level]
        wide = g.pivot(index="query_id", columns="condition", values=r.metric)
        pair = wide[[r.base, r.applied]].dropna()
        zero_count = 0
        if r.metric == "answer_relevancy":
            excluded = pair.eq(0).any(axis=1)
            zero_count = int(excluded.sum())
            pair = pair.loc[~excluded]
        diff = pair[r.applied] - pair[r.base]
        assert len(diff) == r.n and len(g) / 8 == r.n_questions
        assert r.relevancy_zero_excluded == zero_count
        assert r.missing_pairs == r.n_questions - r.n
        for actual, expected in [
            (r.base_mean, pair[r.base].mean()),
            (r.applied_mean, pair[r.applied].mean()),
            (r.mean_delta, diff.mean()),
        ]:
            assert abs(actual - expected) < 1e-10
        if r.metric in QUALITY:
            diff = (
                np.rint(pair[r.applied] * 10000) - np.rint(pair[r.base] * 10000)
            ) / 10000
        threshold = 1 if r.metric == "duration_seconds" else 0.05
        idx = summary.index[
            (summary.grouping == r.grouping)
            & (summary.comparison == r.comparison)
            & (summary.metric == r.metric)
            & (summary.level == r.level)
        ][0]
        stored = summary.loc[idx]
        assert stored["improved"] == int((diff > threshold).sum())
        assert stored["worse"] == int((diff < -threshold).sum())
        assert stored["similar"] == int((diff.abs() <= threshold).sum())
        assert r.n == stored["improved"] + stored["worse"] + stored["similar"]


def validate_tables():
    definitions = json.loads(
        (ROOT / "FINALDOCS/DATA/THESIS_TABLES_60Q.json").read_text(encoding="utf-8")
    )
    wb = load_workbook(ROOT / "FINALDOCS/TABLES/TABLES_60Q.xlsx")
    assert wb.sheetnames == [t["sheet"] for t in definitions]
    for t in definitions:
        ws = wb[t["sheet"]]
        actual = [
            [str(c.value) if c.value is not None else "" for c in row]
            for row in ws.iter_rows(min_row=3)
        ]
        assert actual == [
            [str(v) for v in row] for row in t["rows"]
        ], f"XLSX differs from source: {t['sheet']}"
        assert not FORBIDDEN.search(
            t["title"] + json.dumps(t["rows"], ensure_ascii=False)
        )
        for row in ws.iter_rows(min_row=3):
            for c in row:
                assert c.fill.fill_type is None or c.fill.fgColor.rgb in (
                    "FFFFFFFF",
                    "00FFFFFF",
                ), "Colour shading returned"
        expected_bold = set()
        for scope in t.get("bold_scopes", []):
            col = scope["column"]
            values = {
                row: float(t["rows"][row][col].split()[0].removesuffix("초"))
                for row in scope["rows"]
            }
            maximum = max(values.values())
            expected_bold.update(
                (row, col) for row, value in values.items() if value == maximum
            )
        assert expected_bold == {
            tuple(cell) for cell in t.get("bold_cells", [])
        }, "Maximum bold selection differs"
        for ri, row in enumerate(t["rows"]):
            for ci, value in enumerate(row):
                if value == "":
                    continue  # merged placeholder cell
                assert bool(ws.cell(ri + 3, ci + 1).font.bold) == (
                    ri in t.get("header_rows", [0]) or (ri, ci) in expected_bold
                ), "Unexpected XLSX bold cell"
    assert wb["Appendix_Queries"].max_row == 63
    return len(definitions)


def validate_condition_summary(obs, ann):
    """Recompute every condition mean and SD, overall and per question group."""
    cs = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/CONDITION_SUMMARY.csv", encoding="utf-8-sig"
    )
    raw = obs.merge(ann[["query_id", "question_group"]], on="query_id")
    assert len(cs) == (1 + len(GROUP_SIZES)) * len(CONDITIONS) * len(METRICS)
    for r in cs.itertuples():
        g = raw if r.level == "전체" else raw[raw.question_group == r.level]
        v = g[g.condition == r.condition][r.metric].dropna()
        assert len(v) == r.n
        assert abs(v.mean() - r.mean) < 1e-10, "Condition mean differs"
        assert abs(v.std(ddof=1) - r.sd) < 1e-10, "Condition SD differs"
        assert int((v == 0).sum()) == r.zero_count


def validate_pairs(obs, deltas):
    original = obs.set_index(["query_id", "condition"])
    expected = {
        (q, b, a, m)
        for _, b, a, metrics, _ in comparisons()
        for q in obs.query_id.unique()
        for m in metrics
    }
    actual = set()
    for r in deltas.itertuples():
        key = (r.query_id, r.base, r.applied, r.metric)
        assert key not in actual, "Duplicate paired measurement"
        actual.add(key)
        x, y = (
            original.loc[(r.query_id, r.base), r.metric],
            original.loc[(r.query_id, r.applied), r.metric],
        )
        for v, raw in [(r.base_value, x), (r.applied_value, y)]:
            assert (pd.isna(v) and pd.isna(raw)) or abs(v - raw) < 1e-12
        excluded = r.metric == "answer_relevancy" and (x == 0 or y == 0)
        assert bool(r.relevancy_zero_excluded) == excluded
        if excluded or pd.isna(x) or pd.isna(y):
            assert pd.isna(r.delta), "Excluded or missing pair has a difference"
        else:
            assert abs(r.delta - (y - x)) < 1e-12
    assert actual == expected, "Missing or unapproved comparison"


def validate_cases(records):
    cases = json.loads(
        (ROOT / "FINALDOCS/ANALYSIS/REVIEWED_CASES.json").read_text(encoding="utf-8")
    )
    obs = pd.read_csv(ROOT / "FINALDOCS/DATA/FACTORIAL_OBSERVATIONS_60Q.csv").set_index(
        ["query_id", "condition"]
    )
    for case in cases:
        q = case["query_id"]
        assert case["same_input_passages"] == (
            records[q, case["base"]]["contexts"]
            == records[q, case["applied"]]["contexts"]
        )
        for side in ("base", "applied"):
            rec = records[q, case[side]]
            saved = case[side + "_record"]
            assert saved["answer"] == rec["generated_answer"]
            assert saved["contexts"] == rec["contexts"]
            assert saved["chunk_ids"] == rec["reranked_chunk_ids"]
            for m, scores in case["scores"].items():
                v = obs.loc[(q, case[side]), m]
                assert (pd.isna(v) and scores[side] is None) or abs(
                    v - scores[side]
                ) < 1e-12
        for side, field, needle, presence in case["checks"]:
            saved = case[side + "_record"]
            value = (
                "\n".join(saved["contexts"]) if field == "contexts" else saved["answer"]
            )
            assert (
                needle.lower() in value.lower()
            ) == presence, "Case statement differs from the saved record"
    return cases


def main():
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("--analysis-only", action="store_true")
    args = parser.parse_args()
    obs = pd.read_csv(ROOT / "FINALDOCS/DATA/FACTORIAL_OBSERVATIONS_60Q.csv")
    validate_observations(obs)
    ann = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/QUESTION_CLASSIFICATION_60Q.csv",
        encoding="utf-8-sig",
    )
    assert ann.query_id.is_unique and len(ann) == 60
    assert ann.question_type.value_counts().to_dict() == {
        "방법 및 절차": 29,
        "결과 및 비교": 20,
        "사실 및 정의": 8,
        "목적 및 기여": 3,
    }
    original = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/QUERY_TYPE_ANNOTATIONS.csv", encoding="utf-8-sig"
    ).set_index("query_id")
    records = run_records()
    for r in ann.itertuples():
        assert r.question_type == original.loc[r.query_id, "original_type"]
        assert r.query == records[r.query_id, "H0C0S0"]["query"]
        assert r.required_information and r.type_review and r.answer_span
    d = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/QUERY_LEVEL_DELTAS.csv", encoding="utf-8-sig"
    )
    expected_rows = sum(len(c[3]) for c in comparisons()) * 60
    assert len(d) == expected_rows
    crossing = d.loc[d.metric.isin(QUALITY)].apply(
        lambda r: r.applied[-1] != r.base[-1], axis=1
    )
    assert (
        d.loc[d.metric.isin(QUALITY)][crossing].evaluation_language == "평가 언어 다름"
    ).all() and (
        d.loc[d.metric.isin(QUALITY)][crossing].base == BASE
    ).all(), "Quality comparison crosses evaluation languages without a label"
    assert ann.question_group.value_counts().to_dict() == GROUP_SIZES
    validate_condition_summary(obs, ann)
    for label, base, applied, metrics, language in comparisons():
        g = d[(d.comparison == label) & (d.metric == "korean_ratio")]
        assert len(g) == 60 and set(g.question_type) == set(TYPES)
    summary = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/QUESTION_TYPE_SUMMARY.csv", encoding="utf-8-sig"
    )
    validate_summary(obs, summary, ann)
    assert not any("median" in col for col in summary)
    validate_pairs(obs, d)
    cases = validate_cases(records)
    if args.analysis_only:
        print(
            f"PASS: 60 x 8; 1,920 original quality values; 480 durations; {len(d):,} aligned differences; {len(summary)} summaries; {len(cases)} evidence-checked cases"
        )
        return
    n = validate_tables()
    text = (
        ROOT / "FINALDOCS/MANUSCRIPT/GRADUATION_REPORT_TRANSFER_KO_60Q.md"
    ).read_text(encoding="utf-8")
    assert not FORBIDDEN.search(text), "Unrequested analysis or rejected prose returned"
    assert "# 부록 C" not in text and "fig5_5_question_deltas" not in text
    for old in [
        "T5-1_Mean_SD",
        "T5-2_Runtime",
        "T3-1_Factorial_Design",
        "Appendix_Paired",
        "Appendix_Estimates",
    ]:
        assert old not in text
    figures = re.findall(r"\[그림삽입: ([^|]+)", text)
    result_figs = [f.strip() for f in figures if re.search(r"fig5_[345]_parallel", f)]
    assert len(result_figs) == 3
    for f in figures:
        assert (ROOT / f.strip()).exists()
    for t in json.loads(
        (ROOT / "FINALDOCS/DATA/THESIS_TABLES_60Q.json").read_text(encoding="utf-8")
    ):
        assert text.count("Sheet=" + t["sheet"] + " |") == 1
    cases = validate_cases(records)
    print(
        f"PASS: original 1,920 quality values incl. 5 missing; 480 durations; {len(d):,} differences; three question groups; {n} tables; {len(cases)} intact paired cases"
    )


if __name__ == "__main__":
    main()
