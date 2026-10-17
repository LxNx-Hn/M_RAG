#!/usr/bin/env python3
"""Build a table-safe DOCX transfer copy of the canonical 60-query manuscript.

The Markdown manuscript remains the content authority.  This helper is solely
for importing a structurally sound document into Hancom Office Web: unlike HTML
clipboard transfer, OOXML keeps explicit table-grid widths and embedded figures.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
MANUSCRIPT = ROOT / "FINALDOCS" / "MANUSCRIPT" / "GRADUATION_REPORT_TRANSFER_KO_60Q.md"
WORKBOOK = ROOT / "FINALDOCS" / "TABLES" / "TABLES_60Q.xlsx"
STYLE_MARKER = re.compile(r"\s*\[스타일\s*=\s*[^\]]+\]")
TABLE_MARKER = re.compile(r"^\[표삽입: .*?\| Sheet=([^|\]]+)\|.*\]$")
FIGURE_MARKER = re.compile(r"^\[그림삽입: ([^|\]]+)\|.*\]$")
CODE_FENCE = re.compile(r"^```")

# Values already entered in the user's original HWP cover.  The canonical
# manuscript deliberately retains template placeholders, so this transfer copy
# uses the already-confirmed HWP cover values without touching the source text.
REPLACEMENTS = {
    "[학생명 입력]": "문종건",
    "[지도교수 입력]": "김진석",
    "[제출일 입력]": "2026년 10월 00일",
    "[심사위원 입력]": "심사위원",
    "[인준일 입력]": "2026년 00월 00일",
}


def clean(text: str) -> str:
    text = STYLE_MARKER.sub("", text).strip()
    for before, after in REPLACEMENTS.items():
        text = text.replace(before, after)
    return text


def set_run_font(run, size: float, bold: bool = False, name: str = "함초롬바탕") -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top: int = 90, start: int = 90, bottom: int = 90, end: int = 90) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, cm: float) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    width = tc_pr.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        tc_pr.append(width)
    width.set(qn("w:w"), str(int(cm * 567)))
    width.set(qn("w:type"), "dxa")


def set_table_fixed_layout(table) -> None:
    table.autofit = False
    table_pr = table._tbl.tblPr
    layout = table_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        table_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def table_rows_from_sheet(workbook, sheet_name: str) -> list[list[str]]:
    worksheet = workbook[sheet_name.strip()]
    rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
    rows = [row for row in rows if any(value not in (None, "") for value in row)]
    while rows and str(rows[0][0] or "").startswith("표 "):
        rows.pop(0)
    while rows and str(rows[0][0] or "").startswith("출처:"):
        rows.pop(0)
    return [["" if value is None else str(value) for value in row] for row in rows]


def column_widths(rows: list[list[str]], total_cm: float = 15.5) -> list[float]:
    count = max(len(row) for row in rows)
    weights: list[float] = []
    for col in range(count):
        longest = max((len(row[col]) if col < len(row) else 0) for row in rows)
        weights.append(min(max(longest, 6), 40))
    scale = total_cm / sum(weights)
    return [max(1.2, weight * scale) for weight in weights]


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    count = max(len(row) for row in rows)
    normalized = [row + [""] * (count - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=count)
    table.style = "Table Grid"
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_fixed_layout(table)
    widths = column_widths(normalized)
    # Dense appendices are readable at 7.5pt; ordinary tables retain 8.5pt.
    size = 7.5 if len(normalized) > 18 or count >= 6 else 8.5
    for row_index, values in enumerate(normalized):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[col_index])
            set_cell_margins(cell)
            if row_index == 0:
                set_cell_shading(cell, "E7E6E6")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 or len(value) < 26 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(value)
            set_run_font(run, size=size, bold=row_index == 0)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)


def add_paragraph(doc: Document, text: str, style: str = "Body") -> None:
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.55
    paragraph.paragraph_format.first_line_indent = Cm(0.7) if style == "Body" else Cm(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if style == "Body" else WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    set_run_font(run, size=10.0)


def add_equation(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.0, name="Cambria Math")


def add_picture(doc: Document, relative: str) -> None:
    image_path = ROOT / relative.replace("/", "\\")
    if not image_path.exists():
        raise FileNotFoundError(f"Missing figure: {image_path}")
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(7)
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Cm(14.5))


def heading(doc: Document, text: str, level: int) -> None:
    paragraph = doc.add_paragraph()
    if level == 1:
        paragraph.paragraph_format.page_break_before = True
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(16)
        size = 15.0
    elif level == 2:
        paragraph.paragraph_format.space_before = Pt(16)
        paragraph.paragraph_format.space_after = Pt(8)
        size = 12.0
    else:
        paragraph.paragraph_format.space_before = Pt(10)
        paragraph.paragraph_format.space_after = Pt(6)
        size = 11.0
    run = paragraph.add_run(text)
    set_run_font(run, size=size, bold=True)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    normal = doc.styles["Normal"]
    normal.font.name = "함초롬바탕"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "함초롬바탕")
    normal.font.size = Pt(10)
    if "Body" not in [style.name for style in doc.styles]:
        doc.styles.add_style("Body", WD_STYLE_TYPE.PARAGRAPH)


def build(output: Path) -> None:
    doc = Document()
    configure_document(doc)
    workbook = load_workbook(WORKBOOK, data_only=True)
    lines = MANUSCRIPT.read_text(encoding="utf-8").splitlines()
    in_code = False
    code_lines: list[str] = []
    table_copy_marker = False
    skip_next_tabular_block = False

    for raw in lines:
        stripped = raw.strip()
        if CODE_FENCE.match(stripped):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                is_tabular = any("\t" in item for item in code_lines)
                if is_tabular and skip_next_tabular_block:
                    skip_next_tabular_block = False
                elif is_tabular:
                    add_table(doc, [item.split("\t") for item in code_lines if item.strip()])
                else:
                    for item in code_lines:
                        if item.strip():
                            add_equation(doc, item.strip())
                code_lines = []
                table_copy_marker = False
            continue
        if in_code:
            code_lines.append(raw)
            continue
        if not stripped:
            continue
        if stripped.startswith("[한글 표 복붙용"):
            table_copy_marker = True
            continue
        marker = TABLE_MARKER.match(stripped)
        if marker:
            add_table(doc, table_rows_from_sheet(workbook, marker.group(1)))
            skip_next_tabular_block = True
            continue
        figure = FIGURE_MARKER.match(stripped)
        if figure:
            add_picture(doc, figure.group(1))
            continue
        value = clean(stripped)
        if value.startswith("# "):
            title = value[2:]
            if title == "졸업자격실험보고서":
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_before = Cm(5)
                paragraph.paragraph_format.space_after = Pt(22)
                run = paragraph.add_run(title)
                set_run_font(run, 18.0, bold=True)
            else:
                heading(doc, title, 1)
        elif value.startswith("## "):
            title = value[3:]
            if title.startswith("한국어 질의 기반"):
                paragraph = doc.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.space_after = Cm(10)
                run = paragraph.add_run(title)
                set_run_font(run, 15.0, bold=True)
            else:
                heading(doc, title, 2)
        elif value.startswith("### "):
            heading(doc, value[4:], 3)
        elif value.startswith("[한글 수식 입력기"):
            # The following centered equation is the readable transfer form.
            continue
        else:
            add_paragraph(doc, value)

    # Page-number field in the footer is preserved by Word/Hancom import.
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(args.output)
    print(args.output.resolve())


if __name__ == "__main__":
    main()
