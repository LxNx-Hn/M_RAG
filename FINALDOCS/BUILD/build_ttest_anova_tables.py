"""Rebuild the chapter-5 and appendix statistics tables from TTEST_ANOVA_60Q.json.

Updates FINALDOCS/DATA/THESIS_TABLES_60Q.json and FINALDOCS/TABLES/TABLES_60Q.xlsx
in place, keeping the non-statistical sheets unchanged, and prints each table as
the tab-separated block used in the manuscript.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

ROOT = Path(__file__).resolve().parents[2]
STATS = ROOT / "FINALDOCS/DATA/TTEST_ANOVA_60Q.json"
TABLES_JSON = ROOT / "FINALDOCS/DATA/THESIS_TABLES_60Q.json"
XLSX = ROOT / "FINALDOCS/TABLES/TABLES_60Q.xlsx"
SOURCE_NOTE = "Source: FINALDOCS/DATA/TTEST_ANOVA_60Q.json; paired t-test and RM ANOVA"
RETIRED = [
    "T5-3_Quality_S0",
    "T5-4_Quality_S1",
    "T5-5_Factorial",
    "T5-6_Factorial",
    "T5-7_Paired_Quality",
    "T5-8_Paired_Language",
    "T5-2_Runtime_SD",
]
METRIC_KO = {
    "faithfulness": "근거 충실도",
    "answer_relevancy": "답변 관련성",
    "context_precision": "문맥 정밀도",
    "context_recall": "문맥 재현율",
    "korean_ratio": "한국어 문자 비율",
}
NUMERIC = (
    "n",
    "평균",
    "SD",
    "중앙값",
    "실행 시간",
    "df",
    "F",
    "p",
    "부분",
    "dz",
    "t(",
    "충실도",
    "관련성",
    "정밀도",
    "재현율",
    "한국어 비율",
)


def f3(x: float) -> str:
    return f"{x:.3f}"


def fp(p: float) -> str:
    return "<0.001" if p < 0.001 else f"{p:.3f}"


def msd(mean: float, sd: float) -> str:
    return f"{mean:.3f} ({sd:.3f})"


def paired_row(label, cmp, r, with_means=True):
    row = [label] + ([cmp] if cmp is not None else [])
    if with_means:
        row += [msd(r["on_mean"], r["on_sd"]), msd(r["off_mean"], r["off_sd"])]
    return row + [
        msd(r["mean_diff"], r["sd_diff"]),
        f"{r['t']:.3f} ({r['df']})",
        fp(r["p"]),
        f3(r["dz"]),
    ]


def anova_rows(rows, with_metric=False):
    out = []
    for r in rows:
        row = [METRIC_KO[r["metric"]]] if with_metric else []
        out.append(
            row
            + [
                r["effect"],
                f"{r['df1']}, {r['df2']}",
                f3(r["F"]),
                fp(r["p"]),
                f3(r["partial_eta_squared"]),
            ]
        )
    return out


def build(stats: dict) -> list[dict]:
    d = stats["descriptives"]
    p, s, a, fu = (
        stats["primary_paired_t"],
        stats["secondary_paired_t"],
        stats["rm_anova"],
        stats["interaction_followups"],
    )
    # One results table: S0 rows (English evaluation) above S1 rows (Korean).
    order = sorted(d, key=lambda c: (c[5], c[1], c[3]))
    metrics = [
        "faithfulness",
        "answer_relevancy",
        "context_precision",
        "context_recall",
        "korean_ratio",
        "duration_seconds",
    ]
    header = ["조건", "충실도", "관련성", "정밀도", "재현율", "한국어 비율", "실행 시간(초)"]
    means = [header] + [
        [c] + [msd(d[c][m]["mean"], d[c][m]["sd"]) for m in metrics] for c in order
    ]
    primary = [
        [
            "연구 질문·지표",
            "비교 조건",
            "평균 ON(SD)",
            "평균 OFF(SD)",
            "평균 차이(SD)",
            "t(df)",
            "p",
            "dz",
        ],
        paired_row(
            "RQ1 HyDE 답변 관련성", "H1C0S0−H0C0S0", p["RQ1_HyDE_answer_relevancy"]
        ),
        paired_row("RQ2 CAD 근거 충실도", "H0C1S0−H0C0S0", p["RQ2_CAD_faithfulness"]),
        paired_row(
            "RQ3 SCD 한국어 문자 비율",
            "HyDE OFF, S1−S0",
            p["RQ3_SCD_korean_ratio_H0"],
        ),
    ]
    effect_header = ["효과", "df", "F", "p", "부분 η²"]
    followup = [
        [
            "상호작용과 단순 효과",
            "평균 ON(SD)",
            "평균 OFF(SD)",
            "평균 차이(SD)",
            "t(df)",
            "p",
            "dz",
        ],
        paired_row(
            "한국어 비율 H×C: CAD OFF의 HyDE", None, fu["korean_ratio_HyDE_at_C0"]
        ),
        paired_row(
            "한국어 비율 H×C: CAD ON의 HyDE", None, fu["korean_ratio_HyDE_at_C1"]
        ),
        paired_row("실행 시간 C×S: SCD OFF의 CAD", None, fu["duration_CAD_at_S0"]),
        paired_row("실행 시간 C×S: SCD ON의 CAD", None, fu["duration_CAD_at_S1"]),
    ]
    secondary = [["요인·지표", "평균 차이(SD)", "t(df)", "p", "dz"]]
    for key, label in [
        ("HyDE_faithfulness", "HyDE 근거 충실도"),
        ("HyDE_context_precision", "HyDE 문맥 정밀도"),
        ("HyDE_context_recall", "HyDE 문맥 재현율"),
        ("CAD_answer_relevancy", "CAD 답변 관련성"),
        ("CAD_context_precision", "CAD 문맥 정밀도"),
        ("CAD_context_recall", "CAD 문맥 재현율"),
        ("RQ3_SCD_korean_ratio_all_states", "SCD 한국어 비율(H·C 4상태)"),
    ]:
        secondary.append(paired_row(label, None, s[key], with_means=False))
    quality_header = ["지표", "효과", "df", "F", "p", "부분 η²"]
    return [
        {
            "sheet": "T5-1_Mean_SD",
            "title": "표 5-1 8개 조건의 품질 지표·한국어 비율·실행 시간 평균(SD)",
            "rows": means,
        },
        {
            "sheet": "T5-3_RQ_Paired",
            "title": "표 5-2 연구 질문별 핵심 지표의 ON·OFF 평균(SD)과 대응 t-검정",
            "rows": primary,
        },
        {
            "sheet": "T5-4_Korean_ANOVA",
            "title": "표 5-3 한국어 문자 비율의 H×C×S 반복측정 ANOVA",
            "rows": [effect_header] + anova_rows(a["korean_ratio"]),
        },
        {
            "sheet": "T5-5_Runtime_ANOVA",
            "title": "표 5-4 실행 시간의 H×C×S 반복측정 ANOVA",
            "rows": [effect_header] + anova_rows(a["duration_seconds"]),
        },
        {
            "sheet": "T5-6_Followup",
            "title": "표 5-5 유의한 상호작용(H×C, C×S)의 단순 효과 대응 t-검정",
            "rows": followup,
        },
        {
            "sheet": "T5-7_Secondary",
            "title": "표 5-6 HyDE·CAD 보조 지표와 SCD 한국어 비율의 대응 t-검정",
            "rows": secondary,
        },
        {
            "sheet": "Appendix_Quality_S0",
            "title": "표 B-3 SCD OFF 영어 평가의 품질 지표 H×C 반복측정 ANOVA",
            "rows": [quality_header] + anova_rows(a["quality_S0"], with_metric=True),
        },
        {
            "sheet": "Appendix_Quality_S1",
            "title": "표 B-4 SCD ON 번역 평가의 품질 지표 H×C 반복측정 ANOVA",
            "rows": [quality_header] + anova_rows(a["quality_S1"], with_metric=True),
        },
    ]


ORDER = [
    "T2-1_Related_Work",
    "T3-1_Factorial_Design",
    "T4-1_Environment",
    "T5-1_Mean_SD",
    "T5-3_RQ_Paired",
    "T5-4_Korean_ANOVA",
    "T5-5_Runtime_ANOVA",
    "T5-6_Followup",
    "T5-7_Secondary",
    "Appendix_Queries",
    "Appendix_Paper",
    "Appendix_QueryType",
    "Appendix_Quality_S0",
    "Appendix_Quality_S1",
]


def write_sheet(wb, table):
    if table["sheet"] in wb.sheetnames:
        del wb[table["sheet"]]
    ws = wb.create_sheet(table["sheet"])
    ws.sheet_view.showGridLines = False
    title_font = Font(name="Carlito", size=11)
    ws["A1"], ws["A2"] = table["title"], SOURCE_NOTE
    ws["A1"].font = ws["A2"].font = title_font
    ws.row_dimensions[1].height = 30
    side = Side(style="thin", color="FFD9D9D9")
    border = Border(left=side, right=side, top=side, bottom=side)
    header = table["rows"][0]
    for r, row in enumerate(table["rows"], start=3):
        ws.row_dimensions[r].height = 34
        for c, value in enumerate(row, start=1):
            cell = ws.cell(r, c, value)
            label = str(header[c - 1])
            cell.border = border
            if r == 3:
                cell.font = Font(name="Malgun Gothic", size=10, bold=True)
                cell.fill = PatternFill("solid", fgColor="FFE8E2F4")
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            else:
                cell.font = Font(name="Malgun Gothic", size=10)
                horizontal = (
                    "center"
                    if label in ("조건", "H", "C", "S", "효과")
                    else "right" if label.startswith(NUMERIC) else "left"
                )
                cell.alignment = Alignment(
                    wrap_text=True, vertical="center", horizontal=horizontal
                )
    for c in range(1, len(header) + 1):
        ws.column_dimensions[ws.cell(3, c).column_letter].width = 23
    ws.freeze_panes = "A4"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--blocks", type=Path, help="write manuscript text blocks here")
    args = parser.parse_args()
    stats = json.loads(STATS.read_text(encoding="utf-8"))
    new = {t["sheet"]: t for t in build(stats)}
    existing = {
        t["sheet"]: t for t in json.loads(TABLES_JSON.read_text(encoding="utf-8"))
    }
    for name in RETIRED:
        existing.pop(name, None)
    for name in ("Appendix_Paper", "Appendix_QueryType"):
        rows = existing[name]["rows"]
        if "유효 n" in rows[0]:
            col = rows[0].index("유효 n")
            existing[name]["rows"] = [r[:col] + r[col + 1 :] for r in rows]
        new[name] = existing[name]
    existing.update(new)
    tables = [existing[name] for name in ORDER]
    TABLES_JSON.write_text(
        json.dumps(tables, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    wb = load_workbook(XLSX)
    for name in RETIRED:
        if name in wb.sheetnames:
            del wb[name]
    for table in new.values():
        write_sheet(wb, table)
    wb._sheets = [wb[name] for name in ORDER]
    wb.save(XLSX)
    blocks = "\n\n".join(
        f"## {t['sheet']}\n```text\n"
        + "\n".join("\t".join(str(v) for v in row) for row in t["rows"])
        + "\n```"
        for t in new.values()
    )
    if args.blocks:
        args.blocks.write_text(blocks + "\n", encoding="utf-8")
    print(blocks)


if __name__ == "__main__":
    main()
