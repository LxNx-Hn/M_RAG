"""Build editable OWPML from the school template and immutable thesis inputs.

No Hancom installation, model loading, evaluation API or experiment mutation.
"""

from __future__ import annotations

import argparse
import copy
import hashlib
import json
import re
import shutil
import zipfile
from datetime import datetime, timezone
from pathlib import Path

from hwpx import HwpxDocument
from lxml import etree as ET
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
HH = "http://www.hancom.co.kr/hwpml/2011/head"
NS = {"hp": HP, "hh": HH}


def clean(text: str) -> str:
    text = re.sub(r"\s*\[스타일=[^\]]+\]", "", text)
    text = re.sub(r"^#{1,6}\s+", "", text)
    text = text.replace(" ···· [쪽번호 자동갱신]", "\t")
    return text.replace("`", "").replace("**", "").strip()


def parse_manuscript(path: Path) -> list[dict]:
    """Retain all semantic paragraphs; replace staging blocks with objects."""
    lines = path.read_text(encoding="utf-8").splitlines()
    events = []
    i = 0
    table_block = False
    pending_page_break = False
    while i < len(lines):
        line = lines[i].strip()
        i += 1
        if not line:
            continue
        if line == "[쪽나눔]":
            if pending_page_break:
                raise ValueError("Repeated page-break directive")
            pending_page_break = True
            continue
        if pending_page_break and not line.startswith("## "):
            raise ValueError("Manual page break must precede a section heading")
        if line.startswith("```"):
            block = []
            while i < len(lines) and not lines[i].startswith("```"):
                block.append(lines[i])
                i += 1
            if i == len(lines):
                raise ValueError("Unclosed manuscript code fence")
            i += 1
            if not any("\t" in text for text in block):
                for equation in block:
                    if equation.strip():
                        events.append({"kind": "equation", "text": equation.strip()})
            continue
        match = re.match(r"\[(?:그림|입출력증빙)삽입:\s*([^|]+)\|", line)
        if match:
            pct = re.search(r"본문폭\s*(\d+)%", line)
            events.append(
                {
                    "kind": "picture",
                    "path": match[1].strip(),
                    "fraction": int(pct[1]) / 100 if pct else 1.0,
                }
            )
            table_block = False
            continue
        match = re.match(r"\[표삽입:.*Sheet=([^|]+)\|", line)
        if match:
            events.append({"kind": "table", "sheet": match[1].strip()})
            table_block = True
            continue
        if line.startswith(("[한글 ", "[쪽번호")):
            continue
        # Table staging includes protocol labels already present in XLSX.
        if table_block and line.startswith(("[SCD ", "SCD OFF 평가", "SCD ON 평가")):
            continue
        if not line.startswith("```"):
            table_block = False
        style_match = re.search(r"\[스타일=([^\]]+)\]", line)
        style = style_match[1] if style_match else "본문"
        if re.match(r"^\[\d+\]", line):
            style = "참고문헌리스트"
        events.append(
            {
                "kind": "paragraph",
                "text": clean(line),
                "style": style,
                "page_break": line.startswith("# ") or pending_page_break,
            }
        )
        pending_page_break = False
    if pending_page_break:
        raise ValueError("Dangling page-break directive")
    if sum(e["kind"] == "table" for e in events) != 17:
        raise ValueError("Expected exactly 17 table insertion events")
    if sum(e["kind"] == "picture" for e in events) != 22:
        raise ValueError("Expected exactly 22 picture insertion events")
    return events


def cell_text(cell) -> str:
    value = cell.value
    if value is None:
        return ""  # Empty workbook cell remains empty, never numeric zero.
    if cell.data_type == "f":
        raise ValueError(f"Unresolved spreadsheet formula: {cell.coordinate}")
    if isinstance(value, float):
        fmt = cell.number_format
        if re.fullmatch(r"0\.0+", fmt):
            return f'{value:.{len(fmt.split(".")[1])}f}'
    return str(value)


def table_data(sheet) -> dict:
    rows = [[cell_text(c) for c in row] for row in sheet.iter_rows(min_row=3)]
    merges = [
        [m.min_row - 3, m.min_col - 1, m.max_row - 3, m.max_col - 1]
        for m in sheet.merged_cells.ranges
        if m.min_row >= 3
    ]
    return {"sheet": sheet.title, "rows": rows, "merges": merges}


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def template_audit(doc: HwpxDocument, source: Path) -> dict:
    section = doc.sections[0].element
    secpr = section.find(".//hp:secPr", NS)
    if secpr is None:
        raise ValueError("School template has no section properties")
    page = secpr.find("hp:pagePr", NS)
    margin = page.find("hp:margin", NS)
    col = section.find(".//hp:colPr", NS)
    # Header API is not needed: inspect the converted package via saved XML.
    audit = {
        "source_name": source.name,
        "source_sha256": sha(source),
        "page": dict(page.attrib),
        "margins": dict(margin.attrib),
        "columns": dict(col.attrib),
        "conversion_report": str(doc.conversion_report),
    }
    audit["width"] = (
        int(page.get("width")) - int(margin.get("left")) - int(margin.get("right"))
    )
    return audit


def empty_template(doc: HwpxDocument) -> None:
    """Keep section/column properties and all style definitions, remove examples."""
    section = doc.sections[0]
    secpr = copy.deepcopy(section.element.find(".//hp:secPr", NS))
    colpr = copy.deepcopy(section.element.find(".//hp:colPr", NS))
    for p in list(section.element):
        section.element.remove(p)
    p = doc.add_paragraph("", style="바탕글", inherit_style=False)
    run = p.element.find("hp:run", NS)
    run.insert(0, secpr)
    if colpr is not None:
        ctrl = ET.Element(f"{{{HP}}}ctrl")
        ctrl.append(colpr)
        run.insert(1, ctrl)


def insert_school_covers(doc: HwpxDocument, converted: Path, front: list[dict]) -> None:
    """Reuse the two original cover layout tables, including sizes and gaps."""
    template = HwpxDocument.open(converted)
    source_paragraphs = list(template.sections[0].element)
    covers = [copy.deepcopy(source_paragraphs[i]) for i in (12, 13)]
    first = {
        0: front[0]["text"],
        2: front[1]["text"],
        5: "지도교수: [지도교수 입력]",
        7: "[소속 학과 입력]\n[소속 대학·캠퍼스 입력]",
        9: front[2]["text"],
        11: "[제출연도 입력]",
    }
    second = (
        {
            0: front[3]["text"],
            2: front[4]["text"],
            4: front[5]["text"],
            6: front[6]["text"],
            8: front[7]["text"],
            10: front[8]["text"],
            18: "[소속 대학·학과 입력]",
        }
        if len(front) == 9
        else {}
    )
    for number, (cover, values) in enumerate(zip(covers, (first, second))):
        if number == 1 and not second:
            continue
        cover.set("pageBreak", "0" if number == 0 else "1")
        for cell in cover.findall(".//hp:tc", NS):
            row = int(cell.find("hp:cellAddr", NS).get("rowAddr"))
            sub = cell.find("hp:subList", NS)
            oldp = sub.find("hp:p", NS)
            run = oldp.find("hp:run", NS)
            char_ref = run.get("charPrIDRef") if run is not None else "0"
            attrs = dict(oldp.attrib)
            for child in list(sub):
                sub.remove(child)
            p = ET.SubElement(sub, f"{{{HP}}}p", attrs)
            newrun = ET.SubElement(p, f"{{{HP}}}run", {"charPrIDRef": char_ref})
            ET.SubElement(newrun, f"{{{HP}}}t").text = values.get(row, "")
        doc.sections[0].element.append(cover)


def insert_table(doc: HwpxDocument, data: dict, width: int) -> None:
    rows = data["rows"]
    nr, nc = len(rows), len(rows[0])
    # Wider text columns in appendix keep actual questions and evidence intact.
    if data["sheet"] == "Appendix_Queries":
        fractions = [0.11, 0.29, 0.12, 0.10, 0.05, 0.33]
    elif nc == 2:
        fractions = [0.28, 0.72]
    elif nc == 7:
        fractions = [0.20, 0.11, 0.11, 0.11, 0.22, 0.125, 0.125]
    elif nc == 9:
        fractions = [0.16, 0.11, 0.11, 0.11, 0.11, 0.11, 0.10, 0.10, 0.09]
    else:
        fractions = [1 / nc] * nc
    widths = [int(width * f) for f in fractions]
    widths[-1] += width - sum(widths)
    # Native cell height based on conservative 9 pt CJK width, without cropping.
    heights = []
    for row in rows:
        heights.append(
            max(
                1700,
                max(
                    (
                        1
                        + sum(900 if ord(ch) > 127 else 500 for ch in value)
                        // max(900, w - 1020)
                    )
                    * 1200
                    + 400
                    for value, w in zip(row, widths)
                ),
            )
        )
    table = doc.add_table(
        nr,
        nc,
        width=width,
        height=sum(heights),
        style="본문",
        char_pr_id_ref="31",
        inherit_style=False,
    )
    table.set_column_widths(widths)
    table.element.set("repeatHeader", "1")
    table.element.set("pageBreak", "CELL")
    table.element.find("hp:pos", NS).set("treatAsChar", "0")
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.set_cell_text(r, c, value)
            cell = table.cell(r, c).element
            cell.set("header", "1" if r == 0 else "0")
            cell.find("hp:cellSz", NS).set("height", str(heights[r]))
            for p in cell.findall(".//hp:p", NS):
                p.set("styleIDRef", "13")
                p.set("paraPrIDRef", "18")
            for run in cell.findall(".//hp:run", NS):
                run.set("charPrIDRef", "31")
    for r0, c0, r1, c1 in data["merges"]:
        table.merge_cells(r0, c0, r1, c1)


def emit(
    doc: HwpxDocument, events: list[dict], workbook, width: int, pages: dict
) -> dict:
    mapping = {"paragraphs": [], "tables": [], "pictures": [], "equations": []}
    for e in events:
        kind = e["kind"]
        if kind == "paragraph":
            text = e["text"]
            if e["style"] in ("목차리스트(장)", "목차리스트(절)", "표/그림리스트"):
                label = re.match(r"^\[(?:그림|표) [\dA-Z]+-\d+\]", text)
                key = label[0] if label else text
                if pages:
                    if key not in pages:
                        raise ValueError(f"Unresolved TOC page: {key}")
                    text += "\t" + str(pages[key])
                else:
                    text += "\t"
            p = doc.add_paragraph(text, style=e["style"], inherit_style=False)
            if e["page_break"]:
                p.element.set("pageBreak", "1")
            mapping["paragraphs"].append(e)
        elif kind == "table":
            data = table_data(workbook[e["sheet"]])
            insert_table(doc, data, width)
            mapping["tables"].append(data)
        elif kind == "picture":
            path = ROOT / e["path"]
            with Image.open(path) as image:
                iw, ih = image.size
            w = int(width * e["fraction"])
            h = round(w * ih / iw)
            # Reserve room for caption; original aspect ratio is preserved.
            if h > 62000:
                w = round(w * 62000 / h)
                h = 62000
            doc.add_picture(
                path.read_bytes(),
                "png",
                width=w,
                height=h,
                align="CENTER",
                style="바탕글",
                inherit_style=False,
            )
            mapping["pictures"].append(
                {**e, "sha256": sha(path), "pixels": [iw, ih], "size": [w, h]}
            )
        elif kind == "equation":
            doc.shapes.add_equation(e["text"], base_unit=1000, size=(width, 2600))
            mapping["equations"].append(e["text"])
    doc.page.set_page_number(
        position="BOTTOM_CENTER", format_type="DIGIT", prefix="- ", suffix=" -"
    )
    return mapping


def preserve(path: Path) -> None:
    if path.exists():
        archive = (
            path.parent
            / "versions"
            / datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
        )
        archive.mkdir(parents=True)
        shutil.copy2(path, archive / path.name)


def register_version_part(path: Path) -> None:
    """The HWP converter omits the explicit version.xml manifest item."""
    with zipfile.ZipFile(path) as package:
        parts = {name: package.read(name) for name in package.namelist()}
    # Refresh the file-browser preview; the school's original preview contains
    # style instructions even after its body is replaced with the manuscript.
    parts["Preview/PrvText.txt"] = (
        HwpxDocument.open(path).text.plain()[:4096].encode("utf-8")
    )
    root = ET.fromstring(parts["Contents/content.hpf"])
    manifest = root.find("{http://www.idpf.org/2007/opf/}manifest")
    if not any(item.get("href") == "version.xml" for item in manifest):
        ET.SubElement(
            manifest,
            "{http://www.idpf.org/2007/opf/}item",
            {"id": "version", "href": "version.xml", "media-type": "application/xml"},
        )
    parts["Contents/content.hpf"] = ET.tostring(
        root, encoding="UTF-8", xml_declaration=True
    )
    temporary = path.with_suffix(".building")
    with zipfile.ZipFile(temporary, "w") as package:
        for name, data in parts.items():
            package.writestr(
                name,
                data,
                compress_type=(
                    zipfile.ZIP_STORED if name == "mimetype" else zipfile.ZIP_DEFLATED
                ),
            )
    temporary.replace(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument(
        "--manuscript",
        type=Path,
        default=ROOT / "FINALDOCS/MANUSCRIPT/GRADUATION_REPORT_TRANSFER_KO_60Q.md",
    )
    parser.add_argument(
        "--tables", type=Path, default=ROOT / "FINALDOCS/TABLES/TABLES_60Q.xlsx"
    )
    parser.add_argument("--output", type=Path, default=ROOT / "FINALDOCS/DELIVERY")
    parser.add_argument("--smoke-only", action="store_true")
    parser.add_argument(
        "--page-map",
        type=Path,
        help="Visually reviewed Hancom PDF page map; fails on stale manuscript hash",
    )
    args = parser.parse_args()
    out = args.output
    out.mkdir(parents=True, exist_ok=True)
    doc = HwpxDocument.open(args.template)
    audit = template_audit(doc, args.template)
    if (
        "unconverted=mappingproxy({}), dropped=mappingproxy({})"
        not in str(doc.conversion_report)
        and args.template.suffix.lower() == ".hwp"
    ):
        raise ValueError(f"Template conversion needs review: {doc.conversion_report}")
    template_copy = out / "SCHOOL_TEMPLATE_CONVERTED.hwpx"
    if not template_copy.exists():
        doc.save_to_path(template_copy)
    workbook = load_workbook(args.tables, data_only=False)
    events = parse_manuscript(args.manuscript)
    pages = {}
    if args.page_map:
        page_map = json.loads(args.page_map.read_text(encoding="utf-8"))
        raw = args.manuscript.read_bytes()
        exact_match = page_map["manuscript_sha256"] == sha(args.manuscript)
        lf_match = (
            page_map.get("manuscript_lf_sha256")
            == hashlib.sha256(raw.replace(b"\r\n", b"\n")).hexdigest()
        )
        if not exact_match and not lf_match:
            raise ValueError(
                "Stale page map: export and inspect a new PDF after manuscript edits"
            )
        if page_map["unresolved"]:
            raise ValueError(
                "Page map contains unresolved entries requiring visual review"
            )
        pages = page_map["pages"]
    if args.smoke_only:
        events = [
            events[0],
            events[1],
            events[2],
            {
                "kind": "paragraph",
                "text": "1. 호환성 시험",
                "style": "장(1.)",
                "page_break": True,
            },
            {
                "kind": "paragraph",
                "text": "1.1 편집 가능한 개체",
                "style": "절(1.1)",
                "page_break": False,
            },
            {
                "kind": "paragraph",
                "text": "학교 양식의 문단 스타일과 한글 문장을 검사한다.",
                "style": "본문",
                "page_break": False,
            },
            {
                "kind": "paragraph",
                "text": "표 셀, 그림, 수식과 쪽번호를 각각 확인한다.",
                "style": "본문",
                "page_break": False,
            },
            {"kind": "table", "sheet": "T3-3_Factor_Position"},
            next(e for e in events if e["kind"] == "picture"),
            {"kind": "equation", "text": "z_CAD = (1 + alpha) z_ctx - alpha z_noctx"},
        ]
    empty_template(doc)
    if args.smoke_only:
        front = events[:3]
        insert_school_covers(doc, template_copy, front)
        body_events = events[3:]
    else:
        abstract_index = next(
            i for i, e in enumerate(events) if e.get("text") == "국문초록"
        )
        front = events[:abstract_index]
        if len(front) != 9:
            raise ValueError(
                "Review school-cover mapping: expected nine canonical front paragraphs"
            )
        insert_school_covers(doc, template_copy, front)
        body_events = events[abstract_index:]
    mapping = emit(doc, body_events, workbook, audit["width"], pages)
    mapping["cover_paragraphs"] = front
    mapping["cover_layout_tables"] = 1 if args.smoke_only else 2
    mapping["toc_page_map"] = pages
    name = (
        "HWPX_COMPATIBILITY_SMOKE" if args.smoke_only else "GRADUATION_REPORT_60Q_FINAL"
    )
    target = out / f"{name}.hwpx"
    preserve(target)
    doc.save_to_path(target)
    register_version_part(target)
    reopened = HwpxDocument.open(target)
    validation = reopened.validate()
    if validation.issues:
        raise ValueError(str(validation))
    manifest = {
        **mapping,
        "template": audit,
        "inputs": {
            "manuscript_sha256": sha(args.manuscript),
            "tables_sha256": sha(args.tables),
        },
        "output": {
            "filename": target.name,
            "sha256": sha(target),
            "bytes": target.stat().st_size,
        },
        "xml_validation": str(validation),
    }
    (out / f"{name}.build.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        f'PASS: {target.name}; paragraphs={len(mapping["paragraphs"])}, tables={len(mapping["tables"])}, pictures={len(mapping["pictures"])}, equations={len(mapping["equations"])}'
    )


if __name__ == "__main__":
    main()
