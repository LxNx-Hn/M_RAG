"""Independently check saved HWPX content, objects and original evidence."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import zipfile
from collections import Counter
from pathlib import Path

from hwpx import HwpxDocument
from hwpx.tools.package_validator import validate_package
from layout_checks import validate_layout
from lxml import etree as ET
from openpyxl import load_workbook
from PIL import Image

ROOT = Path(__file__).resolve().parents[2]
HP = "http://www.hancom.co.kr/hwpml/2011/paragraph"
NS = {"hp": HP}


def norm(value: str) -> str:
    return re.sub(r"\s+", "", value)


def canonical_paragraphs(source: str) -> list[str]:
    # Independent of the builder: fenced code is checked separately as table
    # cells/equations; insertion directives and formatting guides are metadata.
    source = re.sub(r"```[^\n]*\n.*?```", "", source, flags=re.DOTALL)
    paragraphs = []
    for line in source.splitlines():
        line = line.strip()
        if not line or re.match(
            r"^\[(?:한글 |그림삽입:|입출력증빙삽입:|표삽입:|쪽나눔\])", line
        ):
            continue
        line = re.sub(r"\s*\[스타일=[^\]]+\]", "", line)
        line = re.sub(r"^#{1,6}\s+", "", line)
        if re.match(r"^\[\d+\] ", line):
            line = re.sub(r"\*([^*]+)\*", r"\1", line)
        line = (
            line.replace(" ···· [쪽번호 자동갱신]", "")
            .replace("`", "")
            .replace("**", "")
        )
        paragraphs.append(line.strip())
    return paragraphs


def node_text(node) -> str:
    return "".join(node.itertext())


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--file",
        type=Path,
        default=ROOT / "FINALDOCS/DELIVERY/GRADUATION_REPORT_60Q_FINAL.hwpx",
    )
    parser.add_argument(
        "--report-dir", type=Path, default=ROOT / "tmp/finaldocs/hwpx_checks"
    )
    args = parser.parse_args()
    delivery = args.file.parent
    args.report_dir.mkdir(parents=True, exist_ok=True)
    source = (
        ROOT / "FINALDOCS/MANUSCRIPT/GRADUATION_REPORT_TRANSFER_KO_60Q.md"
    ).read_text(encoding="utf-8")
    package = validate_package(args.file)
    assert not package.issues, package
    document = HwpxDocument.open(args.file)
    assert not document.validate().issues
    extracted = document.text.plain()
    assert "\ufffd" not in extracted, "Unicode replacement character detected"
    prose_review = json.loads(
        (ROOT / "FINALDOCS/VALIDATION/TEXT_REVIEW.json").read_text(encoding="utf-8")
    )
    for sentence in prose_review["retired_defensive_sentences"]:
        assert sentence not in extracted, "Retired defensive sentence remains in HWPX"
    metadata_text = extracted.replace("FINALDOCS/DATA/evidence_manifest_60q.json", "")
    for marker in (
        "FINALDOCS/",
        "[스타일=",
        "[그림삽입:",
        "[표삽입:",
        "복붙용",
        "쪽번호 자동갱신",
        "[쪽나눔]",
    ):
        assert marker not in metadata_text, f"Unresolved staging metadata: {marker}"
    # The nine front staging entries are mapped into the original school covers.
    # They are metadata, not generic prose to print in the submission form.
    paragraphs = canonical_paragraphs(source)[9:]
    actual_normal = norm(extracted)
    position = 0
    for paragraph in paragraphs:
        needle = norm(paragraph)
        found = actual_normal.find(needle, position)
        assert (
            found >= 0
        ), f"Missing or reordered manuscript paragraph: {paragraph[:100]}"
        position = found + len(needle)
    with zipfile.ZipFile(args.file) as z:
        assert z.testzip() is None
        assert z.infolist()[0].filename == "mimetype"
        assert z.getinfo("mimetype").compress_type == zipfile.ZIP_STORED
        assert z.read("mimetype") == b"application/hwp+zip"
        preview = z.read("Preview/PrvText.txt").decode("utf-8")
        assert "졸업자격실험보고서" in preview and "스타일 적용 방법" not in preview
        for name in z.namelist():
            if name.endswith((".xml", ".hpf", ".rdf")):
                ET.fromstring(z.read(name))
        section = ET.fromstring(z.read("Contents/section0.xml"))
        header = ET.fromstring(z.read("Contents/header.xml"))
        for heading in re.findall(r"\[쪽나눔\]\s*\n## (.*?) \[스타일=", source):
            matches = [
                p
                for p in section.findall("./hp:p", NS)
                if norm("".join(p.xpath("./hp:run/hp:t//text()", namespaces=NS)))
                == norm(heading)
            ]
            assert (
                sum(p.get("pageBreak") == "1" for p in matches) == 1
            ), f"Reviewed section page break missing: {heading}"
        # Exact occurrences catch duplicated prose missed by an order check.
        actual_counts = Counter(
            norm(
                "".join(
                    (
                        (t.text or "")
                        if t.find("hp:tab", NS) is not None
                        else "".join(t.itertext())
                    )
                    for t in p.findall("./hp:run/hp:t", NS)
                )
            )
            for p in section.findall(".//hp:p", NS)
        )
        for paragraph, count in Counter(map(norm, paragraphs)).items():
            assert (
                actual_counts[paragraph] == count
            ), f"Manuscript paragraph occurrence mismatch: {paragraph[:100]}"
        tables = section.findall(".//hp:tbl", NS)
        assert len(tables) >= 12, "Long inline appendix must be segmented"
        workbook = load_workbook(
            ROOT / "FINALDOCS/TABLES/TABLES_60Q.xlsx", data_only=False
        )
        sheets = re.findall(r"\[표삽입:.*?Sheet=([^|]+)\|", source)
        assert len(sheets) == 9
        definitions = {
            t["sheet"]: t
            for t in json.loads(
                (ROOT / "FINALDOCS/DATA/THESIS_TABLES_60Q.json").read_text(
                    encoding="utf-8"
                )
            )
        }
        cell_count = 0
        instances = [
            (t, name, list(range(workbook[name.strip()].max_row - 2)))
            for t, name in zip(tables[2:10], sheets[:8])
        ]
        next_row = 1
        appendix_parts = tables[10:]
        for index, table in enumerate(appendix_parts):
            count = int(table.get("rowCnt"))
            assert count >= 2, "Empty appendix segment"
            source_rows = [0] + list(range(next_row, next_row + count - 1))
            instances.append((table, sheets[-1], source_rows))
            next_row += count - 1
            size = table.find("hp:sz", NS)
            assert int(size.get("height")) <= (
                42000 if index == 0 else 60000
            ), "Inline appendix exceeds page height"
        assert next_row == 61, "Appendix must retain all 60 source rows exactly once"
        continued = [
            "".join(p.xpath("./hp:run/hp:t//text()", namespaces=NS))
            for p in section
            if "".join(p.xpath("./hp:run/hp:t//text()", namespaces=NS)).startswith(
                "[계속표 A-1"
            )
        ]
        assert continued == [
            f"[계속표 A-1 ({i+1}/{len(appendix_parts)})] 60개 질문과 정답 근거"
            for i in range(1, len(appendix_parts))
        ]
        for table, name, source_rows in instances:
            ws = workbook[name.strip()]
            assert int(table.get("rowCnt")) == len(source_rows)
            assert int(table.get("colCnt")) == ws.max_column
            expected_break = "TABLE"
            assert (
                table.get("pageBreak") == expected_break
            ), "Table row pagination changed"
            assert table.get("repeatHeader") == "1"
            merges = {
                (m.min_row - 3, m.min_col - 1): (
                    m.max_row - m.min_row + 1,
                    m.max_col - m.min_col + 1,
                )
                for m in ws.merged_cells.ranges
                if m.min_row >= 3
            }
            for cell in table.findall("hp:tr/hp:tc", NS):
                address = cell.find("hp:cellAddr", NS)
                r, c = int(address.get("rowAddr")), int(address.get("colAddr"))
                original = ws.cell(source_rows[r] + 3, c + 1)
                assert (
                    original.data_type != "f"
                ), "Cached-only formula input not permitted"
                value = "" if original.value is None else str(original.value)
                # This workbook stores thesis precision as text; check every cell.
                if isinstance(original.value, float) and re.fullmatch(
                    r"0\.0+", original.number_format
                ):
                    value = f'{original.value:.{len(original.number_format.split(".")[1])}f}'
                text = "".join(cell.xpath(".//hp:t//text()", namespaces=NS))
                assert norm(text) == norm(
                    value
                ), f"{name} cell {original.coordinate}: {text!r} != {value!r}"
                fill_id = cell.get("borderFillIDRef")
                fill = next(
                    item
                    for item in header.iter()
                    if item.tag.endswith("}borderFill") and item.get("id") == fill_id
                )
                for brush in fill.iter():
                    if brush.tag.endswith("}winBrush"):
                        assert brush.get("faceColor") in (
                            "#FFFFFF",
                            "none",
                            "NONE",
                            None,
                        ), "Result cell colour shading returned"
                for run in cell.findall(".//hp:run", NS):
                    char_id = run.get("charPrIDRef")
                    char = next(
                        item
                        for item in header.iter()
                        if item.tag.endswith("}charPr") and item.get("id") == char_id
                    )
                    is_bold = any(item.tag.endswith("}bold") for item in char)
                    assert is_bold == bool(
                        original.font.bold
                    ), f"Native bold mismatch: {name} {original.coordinate}"
                label = str(ws.cell(3, c + 1).value)
                numeric = (
                    (c in definitions[name.strip()]["numeric_columns"])
                    if definitions[name.strip()].get("numeric_columns") is not None
                    else bool(
                        re.match(
                            r"^(?:n|유효 n|평균|질문 수|기준 점수|적용 점수|기준 시간|적용 시간|점수 차이|개선/|증가/|원문 페이지|근거 충실도|답변 관련성|문맥|한국어 비율|실행 시간)",
                            label,
                        )
                    )
                )
                expected_alignment = (
                    "CENTER"
                    if r == 0
                    or label
                    in (
                        "조건",
                        "H",
                        "C",
                        "S",
                        "효과",
                        "요인",
                        "번호",
                        "HyDE",
                        "CAD",
                        "SCD",
                        "조건 코드",
                        "품질 평가 언어",
                        "비교 조건",
                    )
                    else "RIGHT" if numeric else "LEFT"
                )
                for paragraph in cell.findall(".//hp:p", NS):
                    reference = paragraph.get("paraPrIDRef")
                    shape = next(
                        p
                        for p in header.iter()
                        if p.tag.endswith("}paraPr") and p.get("id") == reference
                    )
                    alignment = next(p for p in shape if p.tag.endswith("}align"))
                    assert (
                        alignment.get("horizontal") == expected_alignment
                    ), f"Table cell alignment changed: {name} {original.coordinate}"
                span = cell.find("hp:cellSpan", NS)
                rs, cs = merges.get((r, c), (1, 1))
                assert (int(span.get("rowSpan")), int(span.get("colSpan"))) == (rs, cs)
                cell_count += 1
        assert workbook["Appendix_Queries"].max_row - 3 == 60
        pictures = section.findall(".//hp:pic", NS)
        image_paths = re.findall(r"\[(?:그림|입출력증빙)삽입:\s*([^|]+)\|", source)
        assert len(pictures) == len(image_paths) == 11
        hpf = ET.fromstring(z.read("Contents/content.hpf"))
        assets = {
            x.get("id"): x.get("href") for x in hpf.iter() if x.tag.endswith("}item")
        }
        for picture, relative in zip(pictures, image_paths):
            ref = picture.find(".//hp:img", NS)
            if ref is None:
                ref = next(x for x in picture.iter() if x.tag.endswith("}img"))
            asset = assets[ref.get("binaryItemIDRef")]
            stored = z.read(asset)
            assert stored == (ROOT / relative.strip()).read_bytes(), relative
            with Image.open(io.BytesIO(stored)) as image:
                iw, ih = image.size
            size = picture.find("hp:sz", NS)
            w, h = int(size.get("width")), int(size.get("height"))
            assert abs(w / h - iw / ih) < 0.001
        equations = section.findall(".//hp:equation", NS)
        scripts = [x.find("hp:script", NS).text for x in equations]
        equation_blocks = [
            block.strip()
            for block in re.findall(r"```text\n(.*?)```", source, re.DOTALL)
            if "\t" not in block
        ]
        assert scripts == equation_blocks and len(scripts) == 8
        template = zipfile.ZipFile(delivery / "SCHOOL_TEMPLATE_CONVERTED.hwpx")
        original_header = ET.fromstring(template.read("Contents/header.xml"))
        layout_report = validate_layout(header, section)
        for tag in ("style", "charPr", "paraPr", "font"):
            originals = original_header.xpath(f'.//*[local-name()="{tag}"]')
            current = header.xpath(f'.//*[local-name()="{tag}"]')
            for orig, now in zip(originals, current):
                assert ET.tostring(orig, method="c14n") == ET.tostring(
                    now, method="c14n"
                ), f"School {tag} modified"
            assert len(current) >= len(originals)
        original_section = ET.fromstring(template.read("Contents/section0.xml"))
        for tag in ("pagePr", "colPr"):
            assert ET.tostring(
                section.find(f".//hp:{tag}", NS), method="c14n"
            ) == ET.tostring(original_section.find(f".//hp:{tag}", NS), method="c14n")
        introduction = next(
            p
            for p in section.findall("./hp:p", NS)
            if p.get("styleIDRef")
            == header.find(
                ".//{http://www.hancom.co.kr/hwpml/2011/head}style[@name='장(1.)']"
            ).get("id")
            if norm("".join(p.xpath("./hp:run/hp:t//text()", namespaces=NS)))
            == norm("1. 서론")
        )
        restart = introduction.find(".//hp:newNum", NS)
        assert (
            restart is not None
            and restart.get("numType") == "PAGE"
            and restart.get("num") == "1"
        ), "Introduction must restart printed page numbering at 1"
        seen = []
        body = source.split("# 1. 서론", 1)[1].split("# 참고문헌", 1)[0]
        for group in re.findall(r"\[(\d+(?:\s*,\s*\d+)*)\]", body):
            for number in map(int, group.split(",")):
                if number not in seen:
                    seen.append(number)
        assert seen == list(range(1, 23)), "References must follow first citation order"
        cover_text = "".join(tables[0].itertext()) + "".join(tables[1].itertext())
        assert "제출 페이지" not in cover_text and "인준 페이지" not in cover_text
        assert (
            "Combination Experiments of HyDE, CAD, and SCD in RAG for Searching English Academic and Technical Documents with Korean Queries"
            in cover_text
        )
        assert "주심" in cover_text and "부심" in cover_text
        template.close()
        assert not re.search(
            r"^# (국문초록|Abstract)|^(주제어|Keywords):", source, re.MULTILINE
        ), "Thesis abstracts returned"
        refs = re.findall(r"^\[(\d+)\]", source.split("# 참고문헌", 1)[1], re.MULTILINE)
        assert refs == [str(n) for n in range(1, 23)]
        for kind, count in (("그림", 11), ("표", 9)):
            captions = re.findall(
                rf"^\[{kind} ([\dA-Z]+-\d+)\].*?\[스타일={kind}제목\]",
                source,
                re.MULTILINE,
            )
            assert len(captions) == len(set(captions)) == count
        assert re.search(
            r"\[그림삽입:[^\n]*fig5_2_io_example[^\n]*\]\s*\[그림 5-2\]", source
        )
    manifest = json.loads(
        (ROOT / "FINALDOCS/DATA/evidence_manifest_60q.json").read_text(encoding="utf-8")
    )
    records = []
    for path, expected in manifest["sources"].items():
        # The immutable manifest contains Windows separators; resolve the same
        # artifact on Linux runners without changing the manifest or its hashes.
        data = (ROOT / path.replace("\\", "/")).read_bytes()
        assert (
            hashlib.sha256(data).hexdigest() == expected
            or hashlib.sha256(data.replace(b"\r\n", b"\n")).hexdigest() == expected
        )
        if path.endswith(".jsonl"):
            records.extend(
                json.loads(line)
                for line in data.decode("utf-8").splitlines()
                if line.strip()
            )
    assert len(records) == 480
    queries = {r["query_id"] for r in records}
    config_key = "config_name" if "config_name" in records[0] else "configuration"
    configs = Counter(r[config_key] for r in records)
    assert len(queries) == 60 and len(configs) == 8 and set(configs.values()) == {60}
    assert len({(r["query_id"], r[config_key]) for r in records}) == 480
    report = {
        "file": args.file.name,
        "sha256": hashlib.sha256(args.file.read_bytes()).hexdigest(),
        "paragraphs_in_order": len(paragraphs),
        "table_cells_verified": cell_count,
        "thesis_tables": 9,
        "native_table_segments": len(tables) - 2,
        "appendix_segments": len(appendix_parts),
        "cover_layout_tables": 2,
        "pictures": 11,
        "editable_equations": 8,
        "references": 22,
        "queries": len(queries),
        "configurations": dict(configs),
        "generation_records": 480,
        "school_styles_preserved": True,
        "package_and_xml": "PASS",
        "layout": layout_report,
        "web_rendering": "see ../VALIDATION/CURRENT_AUDIT.md and PROFESSOR_FEEDBACK.md",
    }
    (args.report_dir / "HWPX_STRUCTURAL_QA.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.report_dir / "HWPX_EXTRACTED_TEXT.txt").write_text(
        "\n".join(line.rstrip() for line in extracted.splitlines()) + "\n",
        encoding="utf-8",
    )
    print(
        f"PASS: HWPX ZIP/XML, {len(paragraphs)} ordered paragraphs, {cell_count} XLSX cells, {len(tables) - 2} native segments of 9 logical tables, 11 pictures, 8 equations, 22 references, 60 x 8 = 480 records, school styles/page/columns"
    )


if __name__ == "__main__":
    main()
