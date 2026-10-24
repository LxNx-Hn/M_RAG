"""Create the thesis tables from the completed question analysis.

Table 3-1 lists the eight conditions with each technique's state. Table 5-1
gives the raw-score mean (SD) of every metric for each condition. Table 5-2
repeats the same means for the three question groups, one column per group
under each metric. Appendix A maps every question to its group and evidence.
The JSON definition and the XLSX workbook are written together; the HWPX
builder reads the workbook.
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "FINALDOCS/ANALYSIS"))
from analyze_questions import (  # noqa: E402
    CONDITIONS,
    CONTEXT_METRICS,
    GROUP_SIZES,
    GROUPS,
    METRICS,
    NAMES,
    NUMBERS,
)

DOC_SHORT = {
    "Mi:dm K 2.5 Pro Technical Report": "Mi:dm",
    "RAG Survey": "RAG Survey",
    "CAD": "CAD",
    "RAPTOR": "RAPTOR",
}
HEADERS = {
    "faithfulness": "근거 충실도",
    "answer_relevancy": "답변 관련성",
    "context_precision": "문맥 정밀도",
    "context_recall": "문맥 재현율",
    "korean_ratio": "한국어 문자 비율",
    "duration_seconds": "처리 시간(초)",
}
GROUP_HEADERS = {
    "수치 질문": "수치\n(18개)",
    "방법 질문": "방법\n(25개)",
    "개념과 결과 질문": "개념과 결과\n(17개)",
}


def condition_label(code):
    return f"{NUMBERS[code]} {NAMES[code]}"


def first_number(cell):
    return float(str(cell).split()[0])


def bold_maxima(rows, row_indexes, col):
    values = {r: first_number(rows[r][col]) for r in row_indexes}
    top = max(values.values())
    return [[r, col] for r, v in values.items() if v == top]


def metric_scopes(body, metric):
    """Rows compared for one column: passage-based metrics inside one evaluation
    language (SCD off rows, SCD on rows), answer-only metrics across all eight,
    and no emphasis for time."""
    if metric == "duration_seconds":
        return []
    if metric in CONTEXT_METRICS:
        return [[r for r, c in body if c[-1] == s] for s in "01"]
    return [[r for r, _ in body]]


def build():
    path = ROOT / "FINALDOCS/DATA/THESIS_TABLES_60Q.json"
    existing = {t["sheet"]: t for t in json.loads(path.read_text(encoding="utf-8"))}
    cs = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/CONDITION_SUMMARY.csv", encoding="utf-8-sig"
    )

    def stat(level, code, metric):
        r = cs[(cs.level == level) & (cs.condition == code) & (cs.metric == metric)]
        assert len(r) == 1
        return r.iloc[0]

    related = existing["T2-1_Related_Work"]
    related["rows"] = [
        ["연구", "핵심 접근", "본 연구와의 관계"],
        ["HyDE[2]", "가상 문서를 임베딩하여 검색", "질문 유형별 답변 관련성과 검색 지표의 변화"],
        ["CAD[3]", "문맥을 포함한 로짓과 제외한 로짓의 결합", "질문 유형별 근거 충실도와 처리 시간의 변화"],
        ["SCD[4]", "언어별 토큰 로짓의 조정", "질문 유형별 한국어 문자 비율의 변화"],
        ["DPR[11], FiD[12]", "밀집 검색과 여러 검색 구간의 활용", "모든 조건에 고정한 검색과 생성 구성"],
        ["Self-RAG[15], Adaptive-RAG[16]", "질문 복잡도에 따른 검색 전략 선택", "질문 유형에 따라 기법을 고르는 기준"],
        ["RAGAS[10], BEIR[19], MTEB[20]", "RAG 자동 평가와 검색 벤치마크", "질문별 원점수와 유형별 평균의 비교"],
    ]
    env = existing["T4-1_Environment"]
    env["title"] = "표 4-1 검색, 재정렬, 생성 단계의 모델과 공통 설정"
    env["rows"] = [
        r
        if r[0] not in ("CAD", "SCD")
        else (
            ["CAD", "문맥 포함 로짓과 제외 로짓의 결합", r[2]]
            if r[0] == "CAD"
            else ["SCD", "목표 언어 토큰의 로짓 조정", r[2]]
        )
        for r in env["rows"]
    ]
    output = [related]

    # Table 3-1: the eight conditions in the order used by every result.
    rows = [["조건", "HyDE", "CAD", "SCD", "품질 평가 근거"]]
    for code in CONDITIONS:
        rows.append(
            [
                condition_label(code),
                "+" if code[1] == "1" else "−",
                "+" if code[3] == "1" else "−",
                "+" if code[5] == "1" else "−",
                "한국어 번역" if code[5] == "1" else "영어",
            ]
        )
    output.append(
        {
            "sheet": "T3-1_Conditions",
            "title": "표 3-1 기법 적용 조건과 품질 평가 근거의 언어",
            "rows": rows,
            "numeric_columns": [],
            "center_columns": [1, 2, 3],
            "column_fractions": [0.3, 0.13, 0.13, 0.13, 0.31],
            "note": "+는 적용, −는 미적용을 뜻한다. 품질 평가 근거는 근거 충실도, 문맥 정밀도, 문맥 재현율을 계산할 때 제공한 검색 결과의 언어이며, 답변 관련성은 질문과 답변만으로 계산한다.",
        }
    )
    output.append(env)

    # Table 5-1: raw-score mean (SD) per condition.
    rows = [["조건"] + [HEADERS[m] for m in METRICS]]
    for code in CONDITIONS:
        cells = [condition_label(code)]
        for m in METRICS:
            s = stat("전체", code, m)
            cells.append(f"{s['mean']:.3f}\n({s['sd']:.3f})")
        rows.append(cells)
    body = list(zip(range(1, len(rows)), CONDITIONS))
    bold, scopes = [], []
    for col, m in enumerate(METRICS, 1):
        for scope in metric_scopes(body, m):
            bold += bold_maxima(rows, scope, col)
            scopes.append({"rows": scope, "column": col})
    output.append(
        {
            "sheet": "T5-1_Condition_Scores",
            "title": "표 5-1 8개 조건의 평가 지표별 평균(표준편차)",
            "rows": rows,
            "bold_cells": bold,
            "bold_scopes": scopes,
            "numeric_columns": [1, 2, 3, 4, 5, 6],
            "column_fractions": [0.22, 0.13, 0.13, 0.13, 0.13, 0.13, 0.13],
            "note": "질문 60개의 평균이며 괄호 안은 표준편차다. 근거 충실도는 결측이 있는 CAD 2개, SCD 3개 질문을 빼고, 처리 시간은 각 생성 실행의 첫 기록인 2개 질문을 빼고 계산하였다. 근거 충실도, 문맥 정밀도, 문맥 재현율은 조건 1, 2, 3, 5를 영어 검색 결과로, 조건 4, 6, 7, 8을 한국어 번역 검색 결과로 평가하였으므로 굵은 값은 각 네 조건 안의 최댓값이다. 답변 관련성과 한국어 문자 비율의 굵은 값은 8개 조건 중 최댓값이다.",
        }
    )

    # Table 5-2: the same means per question group, metrics in three blocks.
    rows, merges, bold, scopes, header_rows = [], [], [], [], []
    for block in [METRICS[0:2], METRICS[2:4], METRICS[4:6]]:
        top = len(rows)
        rows.append(["조건"] + [HEADERS[m] if i == 0 else "" for m in block for i in range(3)])
        rows.append([""] + [GROUP_HEADERS[g] for _ in block for g in GROUPS])
        header_rows += [top, top + 1]
        merges.append([top, 0, top + 1, 0])
        for k in range(len(block)):
            merges.append([top, 1 + 3 * k, top, 3 + 3 * k])
        start = len(rows)
        for code in CONDITIONS:
            cells = [condition_label(code)]
            for m in block:
                for g in GROUPS:
                    cells.append(f"{stat(g, code, m)['mean']:.3f}")
            rows.append(cells)
        body = list(zip(range(start, len(rows)), CONDITIONS))
        for col in range(1, 7):
            for scope in metric_scopes(body, block[(col - 1) // 3]):
                bold += bold_maxima(rows, scope, col)
                scopes.append({"rows": scope, "column": col})
    output.append(
        {
            "sheet": "T5-2_Group_Scores",
            "title": "표 5-2 질문 유형별 8개 조건의 평가 지표 평균",
            "rows": rows,
            "merges": merges,
            "header_rows": header_rows,
            "bold_cells": bold,
            "bold_scopes": scopes,
            "numeric_columns": [1, 2, 3, 4, 5, 6],
            "column_fractions": [0.22, 0.13, 0.13, 0.13, 0.13, 0.13, 0.13],
            "note": "괄호 안은 각 유형의 질문 수다. 근거 충실도는 결측 질문을, 처리 시간은 각 생성 실행의 첫 기록인 2개 질문을 빼고 계산하였다. 근거 충실도, 문맥 정밀도, 문맥 재현율의 굵은 값은 SCD 미적용 네 조건과 적용 네 조건 안에서 각각 구한 최댓값이고, 답변 관련성과 한국어 문자 비율의 굵은 값은 8개 조건 중 최댓값이다.",
        }
    )

    # Appendix A with short document and group labels.
    questions = json.loads(
        (ROOT / "FINALDOCS/DATA/QUESTIONS_60Q.json").read_text(encoding="utf-8")
    )
    groups = pd.read_csv(
        ROOT / "FINALDOCS/ANALYSIS/QUESTION_CLASSIFICATION_60Q.csv",
        encoding="utf-8-sig",
    ).set_index("query_id")["question_group"]
    assert groups.value_counts().to_dict() == GROUP_SIZES
    rows = [["질문 ID", "한국어 질문", "문서", "유형", "쪽", "정답 근거 구간"]]
    for q in questions:
        rows.append(
            [
                q["query_id"],
                q["query"],
                DOC_SHORT[q["document_label"]],
                groups[q["query_id"]].removesuffix(" 질문"),
                str(q["source_page"]),
                q["answer_span"],
            ]
        )
    output.append(
        {
            "sheet": "Appendix_Queries",
            "title": "표 A-1 60개 질문의 대상 문서, 유형과 정답 근거",
            "rows": rows,
            "numeric_columns": [4],
            "column_fractions": [0.165, 0.295, 0.11, 0.12, 0.05, 0.26],
        }
    )
    assert set(GROUPS) == set(GROUP_HEADERS)

    path.write_text(
        json.dumps(output, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    write_xlsx(output)
    contract_path = ROOT / "FINALDOCS/DATA/REVISION_CONTRACT_20261015.json"
    contract = json.loads(contract_path.read_text(encoding="utf-8"))
    contract.update(tables=len(output), sheets=[t["sheet"] for t in output])
    contract_path.write_text(
        json.dumps(contract, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(f"PASS: {len(output)} tables written to JSON and XLSX")
    return output


def write_xlsx(tables):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="777777")
    for t in tables:
        ws = wb.create_sheet(t["sheet"])
        ws.sheet_view.showGridLines = False
        ws.cell(1, 1, t["title"])
        ws.cell(2, 1, "Source: stored 60-question observations; no imputation")
        numeric = set(t.get("numeric_columns", []))
        center = set(t.get("center_columns", []))
        headers = set(t.get("header_rows", [0]))
        bold = {tuple(x) for x in t.get("bold_cells", [])}
        for r, row in enumerate(t["rows"]):
            for c, value in enumerate(row):
                cell = ws.cell(r + 3, c + 1, str(value))
                cell.font = Font(
                    name="Malgun Gothic", size=10, bold=r in headers or (r, c) in bold
                )
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if r in headers or c in center:
                    horizontal = "center"
                else:
                    horizontal = "right" if c in numeric else "left"
                cell.alignment = Alignment(
                    horizontal=horizontal, vertical="center", wrap_text=True
                )
        for r0, c0, r1, c1 in t.get("merges", []):
            ws.merge_cells(
                start_row=r0 + 3, start_column=c0 + 1, end_row=r1 + 3, end_column=c1 + 1
            )
        for c in range(len(t["rows"][0])):
            ws.column_dimensions[chr(65 + c)].width = 20
        ws.freeze_panes = "A4"
    wb.save(ROOT / "FINALDOCS/TABLES/TABLES_60Q.xlsx")


if __name__ == "__main__":
    build()
